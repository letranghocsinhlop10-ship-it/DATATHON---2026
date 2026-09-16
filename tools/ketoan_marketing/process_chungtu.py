#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tu dong doc chung tu chi phi Marketing (Meta Ads + Giay bao no ngan hang + Hoa don
dien tu VN) tu file ZIP/PDF/XML, ghep cap theo So tham chieu va xuat bang hach toan
MISA SME 2023.

Cach dung:
    python3 process_chungtu.py <file.zip|file.pdf|file.xml|thu_muc> ... [-o OUTDIR]

Dau ra trong OUTDIR:
    - bang_hach_toan_misa.csv   (UTF-8 BOM, mo truc tiep bang Excel)
    - bang_hach_toan_misa.xlsx  (neu co openpyxl)
    - bang_hach_toan_misa.md    (bang Markdown de xem nhanh)
    - Bo_Chung_Tu_<SoThamChieu>.pdf (gop cac trang PDF cung mot so tham chieu)
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import shutil
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import datetime

MST_META = "9000000327"
TEN_META = "Meta Platforms Ireland Limited"
MA_DOI_TUONG_META = "FB_META"
TK_NO_CHI_PHI = "6427"
TK_NO_VAT = "1331"
TK_CO_NGAN_HANG = "1121"

THANG_VN = re.compile(r"(\d{1,2})\s*th[aá]ng\s*(\d{1,2}),?\s*(\d{4})", re.IGNORECASE)
REF_META = re.compile(r"S[ốo]\s*tham\s*chi[ếe]u\s*:?\s*([A-Z0-9]{6,24})", re.IGNORECASE)
REF_FACEBK = re.compile(r"FACEBK[\s*/\\-]*([A-Z0-9]{6,24})", re.IGNORECASE)


# --------------------------------------------------------------------------- #
# Tien ich
# --------------------------------------------------------------------------- #
def parse_amount(raw: str | None):
    """'63.337 d' -> 63337.0 ; '97,079 VND' -> 97079.0 ; '1,234.56' -> 1234.56"""
    if raw is None:
        return None
    s = re.sub(r"[^\d.,\-]", "", str(raw))
    if not s or not re.search(r"\d", s):
        return None
    # Nhom 3 chu so deu dan -> dau phan cach hang nghin, bo het
    if re.fullmatch(r"-?\d{1,3}([.,]\d{3})+", s):
        return float(re.sub(r"[.,]", "", s))
    if "." in s and "," in s:
        dec = max(s.rfind("."), s.rfind(","))
        intpart = re.sub(r"[.,]", "", s[:dec])
        return float(f"{intpart}.{s[dec + 1:]}")
    for sep in (".", ","):
        if sep in s:
            head, _, tail = s.rpartition(sep)
            # 1-2 chu so sau dau phan cach -> thap phan; 3 chu so -> hang nghin
            if len(tail) in (1, 2):
                return float(f"{re.sub(r'[.,]', '', head)}.{tail}")
            return float(re.sub(r"[.,]", "", s))
    return float(s)


def fmt_amount(v):
    if v is None:
        return ""
    return f"{v:,.0f}".replace(",", ".") if float(v).is_integer() else f"{v:,.2f}"


def to_ddmmyyyy(d):
    return d.strftime("%d/%m/%Y") if isinstance(d, datetime) else (d or "")


def read_pdf_text(path: str) -> str:
    try:
        import pdfplumber
    except ImportError:
        sys.exit("Thieu thu vien: pip install pdfplumber pypdf openpyxl")
    out = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            out.append(page.extract_text() or "")
    return "\n".join(out)


# --------------------------------------------------------------------------- #
# Mo hinh chung tu
# --------------------------------------------------------------------------- #
@dataclass
class Doc:
    path: str
    kind: str                      # META | BANK | EINVOICE | UNKNOWN
    ref: str | None = None
    date: datetime | None = None
    so_hoa_don: str | None = None
    truoc_thue: float | None = None
    vat: float | None = None
    tong: float | None = None
    thue_suat: str | None = None
    mst_ncc: str | None = None
    ten_ncc: str | None = None
    ma_giao_dich: str | None = None
    ghi_chu: list[str] = field(default_factory=list)
    text: str = ""

    @property
    def name(self):
        return os.path.basename(self.path)


def classify(text: str, path: str) -> str:
    t = text.lower()
    if "meta platforms" in t or "hóa đơn thuế cho" in t or "fbads-" in t:
        return "META"
    if "báo nợ" in t or "debit note" in t or "ghi nợ tài khoản" in t:
        return "BANK"
    if "hóa đơn điện tử" in t or "hóa đơn giá trị gia tăng" in t or path.lower().endswith(".xml"):
        return "EINVOICE"
    return "UNKNOWN"


def find(pattern, text, group=1, flags=re.IGNORECASE):
    m = re.search(pattern, text, flags)
    return m.group(group).strip() if m else None


def parse_meta(doc: Doc):
    t = doc.text
    doc.ref = find(REF_META.pattern, t)
    doc.so_hoa_don = find(r"H[óo]a\s*đ[ơo]n\s*#\s*(\S+)", t)
    doc.mst_ncc = find(r"Tax\s*ID:\s*(\d{9,})", t) or MST_META
    doc.ten_ncc = TEN_META
    doc.ma_giao_dich = find(r"ID giao d[ịi]ch\s*\n?\s*([\d\-]{10,})", t)
    doc.truoc_thue = parse_amount(find(r"T[ổo]ng ph[ụu]\s*:?\s*([\d.,]+)", t))
    m = re.search(r"VAT\s*:?\s*([\d.,]+)\s*[₫đ]?\s*\(Thu[ếe] su[ấa]t:\s*([\d.,]+)\s*%", t, re.IGNORECASE)
    if m:
        doc.vat, doc.thue_suat = parse_amount(m.group(1)), f"{m.group(2)}%"
    else:
        doc.vat = parse_amount(find(r"VAT\s*:?\s*([\d.,]+)", t))
    # Tong tien nam ngay sau so tham chieu tren cung mot dong
    doc.tong = parse_amount(find(r"S[ốo]\s*tham\s*chi[ếe]u\s*:?\s*[A-Z0-9]{6,24}\s+([\d.,]+)\s*[₫đ]", t))
    if doc.tong is None and doc.truoc_thue is not None:
        doc.tong = doc.truoc_thue + (doc.vat or 0)
    m = THANG_VN.search(t)
    if m:
        doc.date = datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))


def parse_bank(doc: Doc):
    t = doc.text
    doc.ref = find(REF_FACEBK.pattern, t)
    doc.ma_giao_dich = find(r"Transaction code:\s*\n?\s*([A-Z0-9\\/_-]{6,})", t)
    doc.tong = parse_amount(find(r"S[ốo] ti[ềe]n/Amount:\s*([\d.,]+)", t))
    d = find(r"(?:Ng[àa]y|Transaction Date)\s*:?\s*(\d{1,2}/\d{1,2}/\d{4})", t)
    if d:
        doc.date = datetime.strptime(d, "%d/%m/%Y")


def parse_einvoice_xml(doc: Doc):
    try:
        root = ET.parse(doc.path).getroot()
    except ET.ParseError as e:
        doc.ghi_chu.append(f"Loi doc XML: {e}")
        return
    g = lambda tag: (root.find(f".//{tag}").text.strip()
                     if root.find(f".//{tag}") is not None and root.find(f".//{tag}").text else None)

    def num(tag):
        """XML hoa don dien tu luon dung dau '.' lam dau thap phan (696.000 = 696 d)."""
        v = g(tag)
        try:
            return float(v) if v is not None else None
        except ValueError:
            return parse_amount(v)
    ky_hieu, so = g("KHHDon"), g("SHDon")
    doc.so_hoa_don = f"{g('KHMSHDon') or ''}{ky_hieu or ''}-{so}" if so else None
    doc.truoc_thue = num("TgTCThue")
    doc.vat = num("TgTThue")
    doc.tong = num("TgTTTBSo")
    doc.thue_suat = g("TSuat")
    nban = root.find(".//NBan")
    if nban is not None:
        doc.ten_ncc = (nban.findtext("Ten") or "").strip() or None
        doc.mst_ncc = (nban.findtext("MST") or "").strip() or None
    nlap = g("NLap")
    if nlap:
        try:
            doc.date = datetime.strptime(nlap[:10], "%Y-%m-%d")
        except ValueError:
            pass


def parse_einvoice_pdf(doc: Doc):
    t = doc.text
    ky_hieu = find(r"K[ýy] hi[ệe]u \(Serial\):\s*(\S+)", t)
    so = find(r"S[ốo] \(No\.\):\s*(\S+)", t)
    doc.so_hoa_don = f"{ky_hieu}-{so}" if so else None
    doc.ten_ncc = find(r"\(Seller\):\s*(.+)", t)
    doc.mst_ncc = find(r"MST \(Tax code\):\s*(\d{9,})", t)
    doc.truoc_thue = parse_amount(find(r"C[ộo]ng ti[ềe]n h[àa]ng \(Subtotal\):\s*([\d.,]+)", t))
    doc.vat = parse_amount(find(r"\(Value added tax\):\s*([\d.,]+)", t))
    doc.tong = parse_amount(find(r"\(Total\):\s*([\d.,]+)", t))
    doc.thue_suat = find(r"Thu[ếe] su[ấa]t \(Tax rate\):\s*([\d.,]+\s*%)", t)
    doc.ma_giao_dich = find(r"S[ốo] tham chi[ếe]u \(Reference\):\s*(\S+)", t)
    d = find(r"Ng[àa]y h[óo]a đ[ơo]n \(Date\):\s*(\d{1,2}/\d{1,2}/\d{4})", t)
    if d:
        doc.date = datetime.strptime(d, "%d/%m/%Y")


def load_doc(path: str) -> Doc:
    if path.lower().endswith(".xml"):
        raw = open(path, encoding="utf-8", errors="ignore").read()
        doc = Doc(path=path, kind="EINVOICE", text=raw)
        parse_einvoice_xml(doc)
    else:
        text = read_pdf_text(path)
        doc = Doc(path=path, kind=classify(text, path), text=text)
        if not text.strip():
            doc.ghi_chu.append("PDF khong co lop text (ban scan) - can OCR")
            return doc
        {"META": parse_meta, "BANK": parse_bank, "EINVOICE": parse_einvoice_pdf}.get(
            doc.kind, lambda d: None
        )(doc)
    # Hoa don dien tu VN mang so tham chieu Meta trong noi dung thanh toan
    if doc.kind == "EINVOICE" and not doc.ref:
        doc.ref = find(REF_FACEBK.pattern, doc.text)
    return doc


def collect_inputs(paths, workdir):
    """Giai nen ZIP (ke ca ZIP long nhau) va tra ve danh sach file PDF/XML."""
    files, queue = [], list(paths)
    while queue:
        p = queue.pop(0)
        if os.path.isdir(p):
            queue += [os.path.join(p, f) for f in sorted(os.listdir(p))]
        elif p.lower().endswith(".zip"):
            dest = tempfile.mkdtemp(dir=workdir)
            with zipfile.ZipFile(p) as z:
                for info in z.infolist():
                    if info.is_dir():
                        continue
                    name = info.filename
                    if info.flag_bits & 0x800 == 0:  # ten file ma CP437 -> tieng Viet
                        try:
                            name = name.encode("cp437").decode("utf-8")
                        except (UnicodeDecodeError, UnicodeEncodeError):
                            pass
                    target = os.path.join(dest, os.path.basename(name))
                    with z.open(info) as src, open(target, "wb") as out:
                        shutil.copyfileobj(src, out)
                    queue.append(target)
        elif p.lower().endswith((".pdf", ".xml")):
            files.append(p)
    return files


# --------------------------------------------------------------------------- #
# Ghep cap + xuat bang
# --------------------------------------------------------------------------- #
def merge_xml_pdf_pairs(docs: list[Doc]) -> list[Doc]:
    """Hoa don dien tu thuong di kem 2 file cung ten goc (.xml du lieu + .pdf ban the hien).
    Gop lai thanh 1 chung tu: lay so lieu tu XML, giu duong dan PDF de gop bo."""
    by_stem: dict[str, list[Doc]] = {}
    for d in docs:
        by_stem.setdefault(os.path.splitext(d.name)[0].lower(), []).append(d)
    out = []
    for stem, items in by_stem.items():
        xml = next((d for d in items if d.path.lower().endswith(".xml")), None)
        pdf = next((d for d in items if d.path.lower().endswith(".pdf")), None)
        if xml and pdf and len(items) == 2:
            for fld in ("ref", "so_hoa_don", "date", "truoc_thue", "vat", "tong",
                        "thue_suat", "mst_ncc", "ten_ncc", "ma_giao_dich"):
                if getattr(xml, fld) in (None, "") and getattr(pdf, fld) not in (None, ""):
                    setattr(xml, fld, getattr(pdf, fld))
            xml.path = pdf.path          # dung ban the hien PDF khi gop bo chung tu
            xml.ghi_chu.append("Da gop du lieu XML + ban the hien PDF")
            out.append(xml)
        else:
            out.extend(items)
    return sorted(out, key=lambda d: d.name)


def build_rows(docs: list[Doc]):
    groups: dict[str, list[Doc]] = {}
    for d in docs:
        groups.setdefault(d.ref or f"(KHONG_CO_REF::{d.name})", []).append(d)

    rows, seq = [], {}
    for ref, items in sorted(groups.items(), key=lambda kv: (
            min((d.date or datetime.max) for d in kv[1]), kv[0])):
        meta = next((d for d in items if d.kind == "META"), None)
        bank = next((d for d in items if d.kind == "BANK"), None)
        einv = next((d for d in items if d.kind == "EINVOICE"), None)

        ngay = (bank.date if bank else None) or (meta.date if meta else None) or (einv.date if einv else None)
        key = ngay.strftime("%Y%m%d") if ngay else "00000000"
        seq[key] = seq.get(key, 0) + 1
        so_ct = f"UNC_{key}_{seq[key]:02d}"

        co, thieu = [], []
        (co if meta else thieu).append("HĐ Meta")
        (co if bank else thieu).append("Báo nợ Bank")
        trang_thai = f"Đủ bộ 2/2 ({'+HĐ GTGT phí NH' if einv else 'Meta+Bank'})" if meta and bank else \
                     f"Thiếu file: {', '.join(thieu)}"

        src = meta or einv or bank
        rows.append({
            "STT": len(rows) + 1,
            "Ngay_Chung_Tu": to_ddmmyyyy(ngay),
            "Ngay_Hach_Toan": to_ddmmyyyy(ngay),
            "So_Chung_Tu": so_ct,
            "Ma_Doi_Tuong": MA_DOI_TUONG_META if meta else (einv.mst_ncc if einv else ""),
            "Ten_Doi_Tuong": TEN_META if meta else ((einv.ten_ncc if einv else "") or ""),
            "Ma_So_Thue_NCC": (meta.mst_ncc if meta else (einv.mst_ncc if einv else "")) or "",
            "Dien_Giai": (f"Chi phí Facebook Ads - HD {meta.so_hoa_don}" if meta
                          else (f"Phí NH/dịch vụ - HD {einv.so_hoa_don}" if einv
                                else f"Thanh toán FACEBK {ref} - chưa có HĐ Meta")),
            "TK_No": TK_NO_CHI_PHI,
            "TK_Co": TK_CO_NGAN_HANG,
            "Tong_Tien": fmt_amount(bank.tong if bank else (src.tong if src else None)),
            "Tien_Truoc_Thue": fmt_amount(src.truoc_thue if src else None),
            "Thue_VAT": fmt_amount(src.vat if src else None),
            "So_Hoa_Don_Meta": (meta.so_hoa_don if meta else (einv.so_hoa_don if einv else "")) or "",
            "So_Tham_Chieu": ref if not ref.startswith("(KHONG_CO_REF") else "",
            "Trang_Thai_Gop": trang_thai,
            "_files": [d.path for d in items],
            "_ref": ref,
        })

        # Hoa don GTGT phi ngan hang di kem la mot but toan doc lap -> tach dong rieng
        if meta and einv:
            seq[key] += 1
            rows.append({
                "STT": len(rows) + 1,
                "Ngay_Chung_Tu": to_ddmmyyyy(einv.date or ngay),
                "Ngay_Hach_Toan": to_ddmmyyyy(einv.date or ngay),
                "So_Chung_Tu": f"UNC_{key}_{seq[key]:02d}",
                "Ma_Doi_Tuong": einv.mst_ncc or "",
                "Ten_Doi_Tuong": einv.ten_ncc or "",
                "Ma_So_Thue_NCC": einv.mst_ncc or "",
                "Dien_Giai": f"Phí dịch vụ ngân hàng GD {ref} - HD {einv.so_hoa_don}",
                "TK_No": f"{TK_NO_CHI_PHI} / {TK_NO_VAT}",
                "TK_Co": TK_CO_NGAN_HANG,
                "Tong_Tien": fmt_amount(einv.tong),
                "Tien_Truoc_Thue": fmt_amount(einv.truoc_thue),
                "Thue_VAT": fmt_amount(einv.vat),
                "So_Hoa_Don_Meta": einv.so_hoa_don or "",
                "So_Tham_Chieu": ref,
                "Trang_Thai_Gop": "HĐ GTGT phí NH (VAT được khấu trừ)",
                "_files": [],
                "_ref": ref,
            })
    return rows


def merge_pdfs(rows, outdir):
    try:
        from pypdf import PdfWriter
    except ImportError:
        return []
    made = []
    for r in rows:
        pdfs = [f for f in r["_files"] if f.lower().endswith(".pdf")]
        if len(pdfs) < 2:
            continue
        ref = r["_ref"]
        out = os.path.join(outdir, f"Bo_Chung_Tu_{ref}.pdf")
        w = PdfWriter()
        for p in pdfs:
            w.append(p)
        with open(out, "wb") as fh:
            w.write(fh)
        made.append(out)
    return made


COLS = ["STT", "Ngay_Chung_Tu", "Ngay_Hach_Toan", "So_Chung_Tu", "Ma_Doi_Tuong",
        "Ten_Doi_Tuong", "Ma_So_Thue_NCC", "Dien_Giai", "TK_No", "TK_Co", "Tong_Tien",
        "Tien_Truoc_Thue", "Thue_VAT", "So_Hoa_Don_Meta", "So_Tham_Chieu", "Trang_Thai_Gop"]


def export(rows, outdir):
    os.makedirs(outdir, exist_ok=True)
    csv_path = os.path.join(outdir, "bang_hach_toan_misa.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    md_path = os.path.join(outdir, "bang_hach_toan_misa.md")
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write("| " + " | ".join(COLS) + " |\n")
        fh.write("|" + "---|" * len(COLS) + "\n")
        for r in rows:
            fh.write("| " + " | ".join(str(r.get(c, "")) for c in COLS) + " |\n")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "MISA_Import"
        ws.append(COLS)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") for c in COLS])
        ws.freeze_panes = "A2"
        for i, c in enumerate(COLS, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(38, len(c) + 6))
        xlsx = os.path.join(outdir, "bang_hach_toan_misa.xlsx")
        wb.save(xlsx)
    except ImportError:
        xlsx = None
    return csv_path, md_path, xlsx


def main():
    ap = argparse.ArgumentParser(description="Doc chung tu Marketing va xuat bang hach toan MISA")
    ap.add_argument("inputs", nargs="+", help="File ZIP/PDF/XML hoac thu muc")
    ap.add_argument("-o", "--outdir", default="ketqua_chungtu")
    args = ap.parse_args()

    workdir = tempfile.mkdtemp(prefix="chungtu_")
    try:
        files = collect_inputs(args.inputs, workdir)
        if not files:
            sys.exit("Khong tim thay file PDF/XML nao trong dau vao.")
        docs = merge_xml_pdf_pairs([load_doc(f) for f in files])

        print(f"Da doc {len(docs)} chung tu:")
        for d in docs:
            print(f"  - [{d.kind:9}] {d.name} | ref={d.ref} | tong={fmt_amount(d.tong)}"
                  + (f" | {'; '.join(d.ghi_chu)}" if d.ghi_chu else ""))

        rows = build_rows(docs)
        csv_p, md_p, xlsx_p = export(rows, args.outdir)
        merged = merge_pdfs(rows, args.outdir)

        print(f"\nDa xuat: {csv_p}\n         {md_p}" + (f"\n         {xlsx_p}" if xlsx_p else ""))
        for m in merged:
            print(f"         {m}")
        for r in rows:
            if r["Trang_Thai_Gop"].startswith("Thiếu"):
                print(f"CANH BAO: {r['So_Tham_Chieu'] or r['_ref']} -> {r['Trang_Thai_Gop']}")
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


if __name__ == "__main__":
    main()
