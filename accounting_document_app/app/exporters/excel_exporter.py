"""Xuất ``ACCOUNTING_RESULT.xlsx`` — 4 sheet theo §22 Phase 1.

    1. HO_SO        một dossier = một dòng
    2. CHUNG_TU     một PDF (chứng từ) = một dòng, KỂ CẢ chưa ghép/lỗi/trùng
    3. MISA_IMPORT  một bút toán = một dòng, sinh bởi ``MisaExporter``
    4. CHECK_ERROR  một vấn đề = một dòng, BLOCKING trước, WARNING/INFO sau

Toàn bộ số tiền ghi dạng số thật (int khi VND nguyên đồng) để Excel tính
tổng được trực tiếp — không bao giờ ghi chuỗi văn bản cho cột tiền.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.config_loader import MisaMappingConfig
from app.exporters.misa_exporter import MisaExporter
from app.models.document import Document
from app.models.dossier import Dossier
from app.models.enums import Severity

__all__ = ["ExportOptions", "ExcelExporter"]

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_SEVERITY_ORDER = {"BLOCKING": 0, "WARNING": 1, "INFO": 2}


@dataclass(frozen=True)
class ExportOptions:
    """Tuỳ chọn cho một lần xuất.

    Attributes:
        voucher_start: Số thứ tự bắt đầu cho Số chứng từ MISA (Q21 — bắt
            buộc, app không tự đoán).
        require_reviewed_for_misa: Chỉ đưa dossier đã review vào sheet MISA.
    """

    voucher_start: int
    require_reviewed_for_misa: bool = False


class ExcelExporter:
    """Dựng file Excel 4 sheet từ dossier + document của một lô.

    Args:
        misa_config: Cấu hình mapping MISA đã nạp.
    """

    def __init__(self, misa_config: MisaMappingConfig) -> None:
        self._misa_config = misa_config
        self._misa_exporter = MisaExporter(misa_config)

    def export(
        self,
        dossiers: list[Dossier],
        documents_by_id: dict[int, Document],
        output_path: Path | str,
        *,
        options: ExportOptions,
    ) -> Path:
        """Xuất workbook 4 sheet ra ``output_path``.

        Args:
            dossiers: Dossier đã validate (bất kể trạng thái — mọi trạng
                thái đều xuất hiện ở sheet HO_SO/CHECK_ERROR để đối chiếu).
            documents_by_id: Tra cứu ``Document`` theo id.
            output_path: Đường dẫn file ``.xlsx`` đích.
            options: Tuỳ chọn xuất (số chứng từ bắt đầu...).

        Returns:
            Đường dẫn file đã ghi.
        """
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        workbook.remove(workbook.active)

        self._write_ho_so(workbook, dossiers, documents_by_id)
        self._write_chung_tu(workbook, dossiers, documents_by_id)
        self._write_misa_import(workbook, dossiers, documents_by_id, options)
        self._write_check_error(workbook, dossiers, documents_by_id)

        workbook.save(output_path)
        logger.info("Đã xuất Excel: %s (%d dossier, %d chứng từ)", output_path, len(dossiers), len(documents_by_id))
        return output_path

    # --------------------------------------------------------------- HO_SO

    def _write_ho_so(
        self, workbook: Workbook, dossiers: list[Dossier], documents_by_id: dict[int, Document]
    ) -> None:
        headers = [
            "HoSo_ID", "Reference", "Status", "Transaction_Date", "Card_Last4",
            "Company_Name", "Tax_Code",
            "Meta_Invoice_Number", "Meta_Transaction_ID", "Meta_Subtotal", "Meta_VAT", "Meta_Total",
            "Debit_Transaction_Code", "Debit_Amount",
            "Bank_Invoice_Serial", "Bank_Invoice_Number", "Bank_Fee_Subtotal", "Bank_Fee_VAT", "Bank_Fee_Total",
            "Folder_Path", "Notes", "Reviewed", "Text_Source",
        ]
        sheet = self._new_sheet(workbook, "HO_SO", headers)

        for dossier in dossiers:
            meta = documents_by_id.get(dossier.meta_document_id)
            debit = documents_by_id.get(dossier.debit_document_id)
            vat = documents_by_id.get(dossier.vat_document_id)
            text_sources = {d.text_source.value for d in (meta, debit, vat) if d is not None}

            sheet.append([
                dossier.dossier_code,
                dossier.reference,
                dossier.status.value,
                dossier.transaction_date,
                dossier.card_last4,
                (meta or debit or vat).value_of("company_name") if (meta or debit or vat) else None,
                (vat or meta).value_of("tax_code") if (vat or meta) else None,
                meta.value_of("invoice_number") if meta else None,
                meta.value_of("transaction_id") if meta else None,
                _num(meta.value_of("subtotal") if meta else None),
                _num(meta.value_of("vat_amount") if meta else None),
                _num(meta.total_amount if meta else None),
                debit.value_of("transaction_code") if debit else None,
                _num(debit.total_amount if debit else None),
                vat.value_of("invoice_serial") if vat else None,
                vat.value_of("invoice_number") if vat else None,
                _num(vat.value_of("subtotal") if vat else None),
                _num(vat.value_of("vat_amount") if vat else None),
                _num(vat.total_amount if vat else None),
                dossier.folder_path,
                dossier.notes,
                "x" if dossier.reviewed else "",
                ",".join(sorted(text_sources)) or None,
            ])

        self._apply_date_column(sheet, headers, "Transaction_Date")
        self._autosize(sheet, headers)

    # ------------------------------------------------------------ CHUNG_TU

    def _write_chung_tu(
        self, workbook: Workbook, dossiers: list[Dossier], documents_by_id: dict[int, Document]
    ) -> None:
        headers = [
            "HoSo_ID", "Document_ID", "Document_Type", "File_Name", "Original_File_Path",
            "Reference", "Invoice_Number", "Invoice_Serial", "Transaction_ID", "Transaction_Code",
            "Document_Date", "Subtotal", "VAT", "Total", "Card_Last4", "Status", "Text_Source", "File_Hash",
        ]
        sheet = self._new_sheet(workbook, "CHUNG_TU", headers)

        document_to_dossier = self._map_document_to_dossier(dossiers)

        for document in documents_by_id.values():
            sheet.append([
                document_to_dossier.get(document.document_id, ""),
                document.document_id,
                document.document_type.value,
                document.file_name,
                str(document.file_path),
                document.match_key,
                document.value_of("invoice_number"),
                document.value_of("invoice_serial"),
                document.value_of("transaction_id"),
                document.value_of("transaction_code"),
                document.document_date,
                _num(document.value_of("subtotal")),
                _num(document.value_of("vat_amount")),
                _num(document.total_amount),
                document.card_last4,
                document.processing_status.value,
                document.text_source.value,
                document.file_hash,
            ])

        self._apply_date_column(sheet, headers, "Document_Date")
        self._autosize(sheet, headers)

    @staticmethod
    def _map_document_to_dossier(dossiers: list[Dossier]) -> dict[int, str]:
        mapping: dict[int, str] = {}
        for dossier in dossiers:
            for document_id in dossier.extra_documents:
                mapping[document_id] = dossier.dossier_code
        return mapping

    # -------------------------------------------------------- MISA_IMPORT

    def _write_misa_import(
        self,
        workbook: Workbook,
        dossiers: list[Dossier],
        documents_by_id: dict[int, Document],
        options: ExportOptions,
    ) -> None:
        headers = [c.header for c in self._misa_config.columns]
        sheet = self._new_sheet(workbook, self._misa_config.sheet_name, headers)

        build = self._misa_exporter.build(
            dossiers,
            documents_by_id,
            voucher_start=options.voucher_start,
            require_reviewed=options.require_reviewed_for_misa,
        )

        for row in build.rows:
            values = []
            for column in self._misa_config.columns:
                values.append(self._misa_cell_value(column.source, row.values))
            sheet.append(values)

        for column in self._misa_config.columns:
            if column.number_format == "dd/mm/yyyy":
                self._apply_date_column(sheet, headers, column.header)

        self._autosize(sheet, headers)

        if build.skipped:
            logger.info("MISA_IMPORT bỏ qua %d dossier: %s", len(build.skipped), build.skipped)

    @staticmethod
    def _misa_cell_value(source: str | None, line_values: dict[str, Any]) -> Any:
        if source is None:
            return None
        _, _, key = source.partition(".")
        value = line_values.get(key)
        if isinstance(value, Decimal):
            return _num(value)
        return value

    # ------------------------------------------------------------ CHECK_ERROR

    def _write_check_error(
        self, workbook: Workbook, dossiers: list[Dossier], documents_by_id: dict[int, Document]
    ) -> None:
        headers = ["HoSo_ID", "Reference", "Severity", "Error_Code", "Description", "File_Name", "Suggested_Action"]
        sheet = self._new_sheet(workbook, "CHECK_ERROR", headers)

        rows: list[tuple[int, list[Any]]] = []
        for dossier in dossiers:
            for issue in dossier.issues:
                document = documents_by_id.get(issue.document_id) if issue.document_id else None
                rows.append((
                    _SEVERITY_ORDER.get(issue.severity.value, 9),
                    [
                        dossier.dossier_code,
                        dossier.reference,
                        issue.severity.value,
                        issue.code,
                        issue.description,
                        document.file_name if document else None,
                        issue.suggested_action,
                    ],
                ))

        for _, row in sorted(rows, key=lambda item: item[0]):
            sheet.append(row)

        self._autosize(sheet, headers)

    # ---------------------------------------------------------------- tiện ích

    @staticmethod
    def _new_sheet(workbook: Workbook, title: str, headers: list[str]) -> Worksheet:
        sheet = workbook.create_sheet(title=title[:31])  # Excel giới hạn 31 ký tự/tên sheet
        sheet.append(headers)
        for col_index in range(1, len(headers) + 1):
            cell = sheet.cell(row=1, column=col_index)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        return sheet

    @staticmethod
    def _apply_date_column(sheet: Worksheet, headers: list[str], header_name: str) -> None:
        if header_name not in headers:
            return
        col_index = headers.index(header_name) + 1
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=col_index)
            if isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"

    @staticmethod
    def _autosize(sheet: Worksheet, headers: list[str], *, max_width: int = 42) -> None:
        for col_index, header in enumerate(headers, start=1):
            width = len(header) + 2
            for row_index in range(2, min(sheet.max_row, 200) + 1):
                value = sheet.cell(row=row_index, column=col_index).value
                if value is not None:
                    width = max(width, len(str(value)) + 2)
            sheet.column_dimensions[get_column_letter(col_index)].width = min(width, max_width)


def _num(value: Decimal | None) -> int | float | None:
    """Chuyển ``Decimal`` sang số Excel thật — int khi là VND nguyên đồng."""
    if value is None:
        return None
    if value == value.to_integral_value():
        return int(value)
    return float(value)
