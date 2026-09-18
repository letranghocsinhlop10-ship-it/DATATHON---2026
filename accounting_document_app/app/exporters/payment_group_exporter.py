"""Xuất Excel cho Payment Case — bank-agnostic (VPBank/VietinBank).

Một sheet riêng (KHÔNG đụng vào ``ExcelExporter``/4 sheet MISA hiện có —
đây là báo cáo bổ sung cho luồng ghép theo payment case, độc lập với luồng
dossier META/DEBIT/VAT gốc). Mỗi ``PaymentCase`` = một dòng, đủ để kế toán
đối chiếu KHÔNG cần mở lại từng PDF.
"""

from __future__ import annotations

import logging
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.matching.payment_group_matcher import bank_doc_txn_id, bank_name_of
from app.models.payment_case import PaymentCase

__all__ = ["PaymentGroupExporter"]

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_SHEET_NAME = "PAYMENT_CASE"

_HEADERS = [
    "Group_Code",
    "Reference",
    "Meta_Transaction_ID",
    "Meta_Invoice_Number",
    "Meta_Paid_Amount",
    "Meta_Payment_Date",
    "Expected_Bank",
    "Bank",
    "Bank_Mismatch",
    "Bank_Transaction_ID",
    "Bank_Main_Amount",
    "Bank_Fee_Amount",
    "Bank_Total_Debit",
    "Bank_Transaction_Date",
    "Card_Last4",
    "VAT_Invoice_Number",
    "VAT_Amount",
    "Meta_Bill_File",
    "Main_Bank_File",
    "Fee_Bank_Files",
    "VAT_Invoice_File",
    "Statement_Source",
    "Statement_Page",
    "Match_Status",
    "Match_Confidence",
    "Match_Reason",
]


def _num(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    return int(value) if value == value.to_integral_value() else float(value)


class PaymentGroupExporter:
    """Dựng workbook 1 sheet từ danh sách ``PaymentCase``."""

    def export(self, groups: list[PaymentCase], output_path: Path | str) -> Path:
        """Xuất ``output_path`` — mỗi case một dòng.

        Args:
            groups: Kết quả từ ``PaymentGroupMatcher.match()``.
            output_path: Đường dẫn file ``.xlsx`` đích.

        Returns:
            ``output_path`` sau khi ghi xong.
        """
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        workbook.remove(workbook.active)
        sheet = self._new_sheet(workbook, _SHEET_NAME, _HEADERS)

        for case in groups:
            sheet.append(self._row(case))

        self._autosize(sheet, _HEADERS)
        workbook.save(output_path)
        logger.info("Đã xuất %s: %s (%d payment case)", _SHEET_NAME, output_path, len(groups))
        return output_path

    @staticmethod
    def _row(case: PaymentCase) -> list:
        bill = case.meta_bill
        main = case.main_payment
        vat = case.vat_invoice
        fee_amount = sum((f.total_amount for f in case.fees if f.total_amount is not None), Decimal(0))
        statement = case.statement_rows[0] if case.statement_rows else None
        main_bank_file = (
            main.file_name if main is not None else ("(trích xuất từ sao kê)" if statement else None)
        )

        return [
            case.group_code,
            case.reference,
            bill.value_of("transaction_id") if bill else None,
            bill.value_of("invoice_number") if bill else None,
            _num(bill.total_amount) if bill else None,
            bill.document_date if bill else None,
            case.expected_bank,
            case.bank_name or (bank_name_of(main) if main else None),
            "x" if case.bank_mismatch else "",
            bank_doc_txn_id(main) if main else None,
            _num(main.total_amount) if main else None,
            _num(fee_amount) if case.fees else None,
            _num(case.bank_total_debit),
            case.transaction_date,
            (main.card_last4 if main else None) or (bill.card_last4 if bill else None),
            vat.value_of("invoice_number") if vat else None,
            _num(vat.total_amount) if vat else None,
            bill.file_name if bill else None,
            main_bank_file,
            "; ".join(f.file_name for f in case.fees) or None,
            vat.file_name if vat else None,
            statement.source_pdf.name if statement and statement.source_pdf else None,
            statement.source_page if statement else None,
            case.status.value,
            case.confidence,
            case.reason,
        ]

    @staticmethod
    def _new_sheet(workbook: Workbook, title: str, headers: list[str]) -> Worksheet:
        sheet = workbook.create_sheet(title=title[:31])
        sheet.append(headers)
        for col_index in range(1, len(headers) + 1):
            cell = sheet.cell(row=1, column=col_index)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        return sheet

    @staticmethod
    def _autosize(sheet: Worksheet, headers: list[str], *, max_width: int = 42) -> None:
        for col_index, header in enumerate(headers, start=1):
            width = len(header) + 2
            for row_index in range(2, min(sheet.max_row, 200) + 1):
                value = sheet.cell(row=row_index, column=col_index).value
                if value is not None:
                    width = max(width, len(str(value)) + 2)
            sheet.column_dimensions[get_column_letter(col_index)].width = min(width, max_width)
