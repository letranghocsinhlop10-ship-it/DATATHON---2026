"""Test xuất Excel 4 sheet — mở lại file thật bằng openpyxl để kiểm tra nội dung."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.config_loader import load_accounting_config, load_misa_mapping_config
from app.exporters.excel_exporter import ExcelExporter, ExportOptions
from app.matching.dossier_builder import DossierBuilder
from app.matching.dossier_validator import DossierValidator
from matching_helpers import Decimal as D
from matching_helpers import date, debit, meta, vat


@pytest.fixture(scope="module")
def misa_config():
    return load_misa_mapping_config(load_accounting_config())


def _build_valid_dossier():
    docs = [
        meta(
            "ABCD1234EF", subtotal=D("1000000"), vat_amount=D("100000"), total_amount=D("1100000"),
            invoice_number="FBADS-100-1", card_last4="1234", document_date=date(2026, 8, 1),
            company_name="CTY ABC",
        ),
        debit("ABCD1234EF", "FT1", total_amount=D("1122000"), card_last4="1234"),
        vat(
            "ABCD1234EF", "FT1", subtotal=D("20000"), vat_amount=D("2000"), total_amount=D("22000"),
            invoice_serial="1K26AAA", invoice_number="00100001", tax_code="0100000001",
        ),
    ]
    result = DossierBuilder().build(docs)
    docs_by_id = {d.document_id: d for d in docs}
    DossierValidator().validate(result.dossiers[0], docs_by_id)
    return result.dossiers, docs_by_id


class TestBonSheet:
    def test_dung_bon_sheet_dung_ten(self, misa_config, tmp_path):
        dossiers, docs_by_id = _build_valid_dossier()
        out = ExcelExporter(misa_config).export(
            dossiers, docs_by_id, tmp_path / "out.xlsx", options=ExportOptions(voucher_start=1)
        )
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["HO_SO", "CHUNG_TU", misa_config.sheet_name, "CHECK_ERROR"]

    def test_ho_so_co_dung_1_dong(self, misa_config, tmp_path):
        dossiers, docs_by_id = _build_valid_dossier()
        out = ExcelExporter(misa_config).export(
            dossiers, docs_by_id, tmp_path / "out.xlsx", options=ExportOptions(voucher_start=1)
        )
        ws = openpyxl.load_workbook(out)["HO_SO"]
        assert ws.max_row == 2  # header + 1 dossier
        header = [c.value for c in ws[1]]
        row = dict(zip(header, next(ws.iter_rows(min_row=2, max_row=2, values_only=True))))
        assert row["HoSo_ID"] == "HS000001"
        assert row["Status"] == "VALID"
        assert row["Meta_Total"] == 1100000
        assert isinstance(row["Meta_Total"], int)

    def test_chung_tu_co_dung_3_dong(self, misa_config, tmp_path):
        dossiers, docs_by_id = _build_valid_dossier()
        out = ExcelExporter(misa_config).export(
            dossiers, docs_by_id, tmp_path / "out.xlsx", options=ExportOptions(voucher_start=1)
        )
        ws = openpyxl.load_workbook(out)["CHUNG_TU"]
        assert ws.max_row == 4  # header + 3 chứng từ
        header = [c.value for c in ws[1]]
        hoso_col = header.index("HoSo_ID")
        for row in ws.iter_rows(min_row=2, values_only=True):
            assert row[hoso_col] == "HS000001"

    def test_misa_sheet_co_dung_4_dong(self, misa_config, tmp_path):
        dossiers, docs_by_id = _build_valid_dossier()
        out = ExcelExporter(misa_config).export(
            dossiers, docs_by_id, tmp_path / "out.xlsx", options=ExportOptions(voucher_start=52601)
        )
        ws = openpyxl.load_workbook(out)[misa_config.sheet_name]
        assert ws.max_row == 5  # header + 4 bút toán
        assert ws.max_column == 34
        header = [c.value for c in ws[1]]
        so_ct_col = header.index("Số chứng từ (*)") + 1
        assert ws.cell(row=2, column=so_ct_col).value == "NVK052601"

    def test_misa_sheet_giu_nguyen_34_cot_ke_ca_cot_trong(self, misa_config, tmp_path):
        dossiers, docs_by_id = _build_valid_dossier()
        out = ExcelExporter(misa_config).export(
            dossiers, docs_by_id, tmp_path / "out.xlsx", options=ExportOptions(voucher_start=1)
        )
        ws = openpyxl.load_workbook(out)[misa_config.sheet_name]
        header = [c.value for c in ws[1]]
        assert header[0] == "Hiển thị trên sổ"
        # Cột "Hiển thị trên sổ" luôn để trống theo file mẫu thật.
        for row in ws.iter_rows(min_row=2, values_only=True):
            assert row[0] is None

    def test_check_error_sap_xep_blocking_truoc(self, misa_config, tmp_path):
        docs = [meta("ABCD1234EF", subtotal=D("1000000"))]  # thiếu debit + vat
        result = DossierBuilder().build(docs)
        docs_by_id = {d.document_id: d for d in docs}
        DossierValidator().validate(result.dossiers[0], docs_by_id)

        out = ExcelExporter(misa_config).export(
            result.dossiers, docs_by_id, tmp_path / "out.xlsx", options=ExportOptions(voucher_start=1)
        )
        ws = openpyxl.load_workbook(out)["CHECK_ERROR"]
        header = [c.value for c in ws[1]]
        sev_col = header.index("Severity")
        severities = [row[sev_col] for row in ws.iter_rows(min_row=2, values_only=True)]
        assert severities[0] == "BLOCKING"


class TestNgayThangDinhDangDung:
    def test_cot_ngay_co_number_format_dd_mm_yyyy(self, misa_config, tmp_path):
        dossiers, docs_by_id = _build_valid_dossier()
        out = ExcelExporter(misa_config).export(
            dossiers, docs_by_id, tmp_path / "out.xlsx", options=ExportOptions(voucher_start=1)
        )
        ws = openpyxl.load_workbook(out)["HO_SO"]
        header = [c.value for c in ws[1]]
        col = header.index("Transaction_Date") + 1
        assert ws.cell(row=2, column=col).number_format == "dd/mm/yyyy"


class TestKhongDocDuocThiDeTrong:
    def test_dossier_thieu_meta_khong_bia_gia_tri(self, misa_config, tmp_path):
        docs = [debit("ABCD1234EF", "FT1"), vat("ABCD1234EF", "FT1")]
        result = DossierBuilder().build(docs)
        docs_by_id = {d.document_id: d for d in docs}
        DossierValidator().validate(result.dossiers[0], docs_by_id)

        out = ExcelExporter(misa_config).export(
            result.dossiers, docs_by_id, tmp_path / "out.xlsx", options=ExportOptions(voucher_start=1)
        )
        ws = openpyxl.load_workbook(out)["HO_SO"]
        header = [c.value for c in ws[1]]
        row = dict(zip(header, next(ws.iter_rows(min_row=2, max_row=2, values_only=True))))
        assert row["Meta_Invoice_Number"] is None
        assert row["Meta_Total"] is None
