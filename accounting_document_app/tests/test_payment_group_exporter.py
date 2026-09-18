"""Test xuất Excel cho Facebook Payment Group (§K)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl

from app.core.bank_statement_parser import BankTransaction
from app.exporters.payment_group_exporter import PaymentGroupExporter
from app.models.enums import DocumentType, PaymentGroupStatus
from app.models.facebook_payment_group import FacebookPaymentGroup
from matching_helpers import make_doc


class TestXuatExcel:
    def test_mot_group_mot_dong(self, tmp_path):
        bill = make_doc(
            DocumentType.META_INVOICE,
            reference_number="ABCD1234EF",
            transaction_id="123-456",
            invoice_number="FBADS-1-1",
            total_amount=Decimal("2530672"),
            document_date=date(2026, 8, 1),
        )
        main = make_doc(
            DocumentType.VPBANK_DEBIT_NOTE,
            meta_reference="ABCD1234EF",
            transaction_code="FT100000001",
            total_amount=Decimal("2558510"),
            card_last4="1234",
        )
        group = FacebookPaymentGroup(
            group_code="FB_ABCD1234EF_2026-08-01",
            facebook_reference="ABCD1234EF",
            transaction_date=date(2026, 8, 1),
            facebook_bill=bill,
            main_payment=main,
            status=PaymentGroupStatus.MATCHED_HIGH,
            confidence="HIGH",
            reason="Facebook reference khớp tuyệt đối: ABCD1234EF",
        )

        out = tmp_path / "payment_groups.xlsx"
        PaymentGroupExporter().export([group], out)

        wb = openpyxl.load_workbook(out)
        sheet = wb["FACEBOOK_PAYMENT_GROUP"]
        header = [c.value for c in sheet[1]]
        row = {h: sheet.cell(row=2, column=i + 1).value for i, h in enumerate(header)}

        assert row["Facebook_Reference"] == "ABCD1234EF"
        assert row["Facebook_Paid_Amount"] == 2530672
        assert row["Bank"] == "VPBANK"
        assert row["Bank_Transaction_ID"] == "FT100000001"
        assert row["Bank_Main_Amount"] == 2558510
        assert row["Card_Last4"] == "1234"
        assert row["Match_Status"] == "MATCHED_HIGH"
        assert "ABCD1234EF" in row["Match_Reason"]

    def test_tong_phi_tinh_dung_khi_co_nhieu_fee(self, tmp_path):
        main = make_doc(DocumentType.VIETINBANK_DEBIT_ADVICE, transaction_number="9001", total_amount=Decimal("100"))
        fee1 = make_doc(DocumentType.VIETINBANK_DEBIT_ADVICE, transaction_number="9000", total_amount=Decimal("10"))
        fee2 = make_doc(DocumentType.VIETINBANK_DEBIT_ADVICE, transaction_number="8999", total_amount=Decimal("5"))
        group = FacebookPaymentGroup(
            group_code="FB_X_2026-08-31", main_payment=main, fees=[fee1, fee2],
            status=PaymentGroupStatus.MATCHED_HIGH,
        )
        out = tmp_path / "payment_groups.xlsx"
        PaymentGroupExporter().export([group], out)

        wb = openpyxl.load_workbook(out)
        sheet = wb["FACEBOOK_PAYMENT_GROUP"]
        header = [c.value for c in sheet[1]]
        row = {h: sheet.cell(row=2, column=i + 1).value for i, h in enumerate(header)}
        assert row["Bank_Fee_Amount"] == 15
        assert row["Bank_Total_Debit"] == 115  # main + hai fee
        assert row["Bank"] == "VIETINBANK"

    def test_khong_co_phieu_ngan_hang_goc_ghi_ro_trich_xuat_tu_sao_ke(self, tmp_path):
        bill = make_doc(DocumentType.META_INVOICE, reference_number="ABCD1234EF")
        row_data = BankTransaction(
            stt=1, transaction_id="FT1", value_date=date(2026, 8, 1), transaction_time="10:00",
            debit_amount=Decimal("1"), credit_amount=None, transaction_detail="x",
            running_balance=Decimal("1"), facebook_reference="ABCD1234EF",
            source_pdf=None, source_page=1,
        )
        group = FacebookPaymentGroup(
            group_code="FB_ABCD1234EF_2026-08-01", facebook_bill=bill, main_payment=None,
            statement_rows=[row_data], status=PaymentGroupStatus.MATCHED,
        )
        out = tmp_path / "payment_groups.xlsx"
        PaymentGroupExporter().export([group], out)

        wb = openpyxl.load_workbook(out)
        sheet = wb["FACEBOOK_PAYMENT_GROUP"]
        header = [c.value for c in sheet[1]]
        row = {h: sheet.cell(row=2, column=i + 1).value for i, h in enumerate(header)}
        assert row["Main_Bank_File"] == "(trích xuất từ sao kê)"
