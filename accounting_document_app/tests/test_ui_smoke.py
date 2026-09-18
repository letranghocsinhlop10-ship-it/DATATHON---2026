"""Smoke test GUI — dựng widget thật qua Qt offscreen, kiểm tra hành vi cơ bản.

Dùng fixture ``qtbot`` của pytest-qt (không tự định nghĩa ``qapp`` — trùng
tên với fixture có sẵn của pytest-qt từng gây treo do mất cơ chế dọn widget
giữa các test). Mọi widget/dialog dựng ra đều đăng ký qua
``qtbot.addWidget()`` để pytest-qt tự đóng/dọn sau mỗi test.

Không gọi ``.exec()`` trên dialog blocking (``QMessageBox``) mà không có ai
bấm nút — trong môi trường offscreen sẽ treo vô hạn. Dùng ``monkeypatch``
để thay các lời gọi ``QMessageBox.*`` bằng hàm không chặn.
"""

from __future__ import annotations

import os
from datetime import date
from pathlib import Path

import pytest
from PySide6.QtWidgets import QMessageBox

from app.database.database import Database
from app.database.document_repository import DocumentRepository
from app.database.dossier_repository import DossierRepository
from app.database.settings_repository import SettingsRepository
from app.models.dossier import Dossier
from app.models.document import Document
from app.models.enums import DocumentType
from app.models.extracted_field import ExtractedField, FieldSet
from app.utils.file_utils import iter_pdf_files

# Repo công khai — KHÔNG hard-code đường dẫn/tên file PDF thật (tên file có
# thể chứa số hoá đơn thật). Test tuỳ chọn dưới đây tự bỏ qua nếu biến môi
# trường này không được đặt, giống cơ chế ở tests/test_real_samples.py.
_SAMPLE_DIR = os.environ.get("ACCOUNTING_SAMPLE_DIR")


def _first_sample_pdf() -> Path | None:
    if not _SAMPLE_DIR:
        return None
    return next(iter_pdf_files(Path(_SAMPLE_DIR)), None)


@pytest.fixture
def db():
    with Database(":memory:") as database:
        yield database


@pytest.fixture(autouse=True)
def _no_blocking_messagebox(monkeypatch):
    """Chặn mọi QMessageBox modal trở thành no-op — tránh treo môi trường
    offscreen không có ai bấm nút. Test nào cần kiểm tra ĐÃ GỌI hộp thoại
    hay chưa thì tự ghi đè thêm bằng monkeypatch riêng của nó."""
    for name in ("information", "warning", "critical"):
        monkeypatch.setattr(QMessageBox, name, staticmethod(lambda *a, **k: None))
    monkeypatch.setattr(
        QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes)
    )


class TestMainWindow:
    def test_dung_thanh_cong(self, qtbot, db, classifier_config, extraction_config, app_settings, accounting_config, misa_config):
        from app.ui.main_window import MainWindow

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)
        assert win.windowTitle() == "Marketing Accounting Document Tool"
        assert win._scan_button.isEnabled() is True

    def test_match_bi_khoa_khi_chua_scan(self, qtbot, db, classifier_config, extraction_config, app_settings, accounting_config, misa_config):
        from app.ui.main_window import MainWindow

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)
        assert win._match_button.isEnabled() is False

    def test_organize_va_export_bi_khoa_luc_dau(
        self, qtbot, db, classifier_config, extraction_config, app_settings, accounting_config, misa_config
    ):
        from app.ui.main_window import MainWindow

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)
        assert win._organize_button.isEnabled() is False
        assert win._export_button.isEnabled() is False

    def test_scan_thu_muc_trong_bao_loi(
        self, qtbot, db, classifier_config, extraction_config, app_settings,
        accounting_config, misa_config, monkeypatch,
    ):
        from app.ui.main_window import MainWindow

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)
        win._input_folder_edit.setText("")  # chưa chọn thư mục

        called = {}
        monkeypatch.setattr(
            QMessageBox, "warning", staticmethod(lambda *a, **k: called.setdefault("warned", True))
        )
        win._on_scan_clicked()
        assert called.get("warned") is True

    def test_populate_bang_dossier(self, qtbot, db, classifier_config, extraction_config, app_settings, accounting_config, misa_config):
        from app.models.enums import DossierStatus
        from app.ui.main_window import MainWindow

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)
        dossier = Dossier(
            dossier_code="HS000001", reference="ABCD1234EF",
            meta_count=1, debit_count=1, vat_count=1, status=DossierStatus.VALID,
        )
        win._dossier_rows = [dossier]
        win._populate_dossier_table([dossier])
        assert win._dossier_table.rowCount() == 1
        assert win._dossier_table.item(0, 0).text() == "HS000001"


class TestLayPdfTuZip:
    """Bấm nút '📦 Lấy PDF từ ZIP' thật — qua ZipExtractWorker (QThread) chứ
    không gọi thẳng ``extract_pdfs_from_zips`` — theo đúng mẫu đã bắt được
    bug sqlite3 check_same_thread ở SCAN/MATCH trước đây."""

    def test_bam_nut_giai_nen_dua_pdf_vao_all_data_va_cap_nhat_thu_muc_nguon(
        self, qtbot, db, classifier_config, extraction_config, app_settings,
        accounting_config, misa_config, monkeypatch, tmp_path,
    ):
        import zipfile

        from PySide6.QtWidgets import QFileDialog

        from app.ui.main_window import MainWindow

        source_dir = tmp_path / "data_tool_read_pdf"
        source_dir.mkdir()
        with zipfile.ZipFile(source_dir / "invoice (1).zip", "w") as zf:
            zf.writestr("invoice.pdf", b"%PDF-1.4 noi dung gia")
            zf.writestr("invoice.xml", b"<xml/>")

        monkeypatch.setattr(
            QFileDialog, "getExistingDirectory", staticmethod(lambda *a, **k: str(source_dir))
        )

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)

        win._on_extract_zip_clicked()
        assert win._zip_extract_worker is not None
        with qtbot.waitSignal(win._zip_extract_worker.finished_ok, timeout=10000):
            pass
        qtbot.waitUntil(lambda: not win._zip_extract_worker.isRunning(), timeout=2000)

        dest = source_dir / "ALL_DATA"
        assert dest.is_dir()
        assert (dest / "invoice.pdf").is_file()
        assert not (dest / "invoice.xml").exists()
        # Tích hợp pipeline: ALL_DATA phải được điền thẳng vào ô "Thư mục PDF".
        assert win._input_folder_edit.text() == str(dest)

    def test_khong_chon_thu_muc_thi_khong_lam_gi(
        self, qtbot, db, classifier_config, extraction_config, app_settings,
        accounting_config, misa_config, monkeypatch,
    ):
        from PySide6.QtWidgets import QFileDialog

        from app.ui.main_window import MainWindow

        monkeypatch.setattr(QFileDialog, "getExistingDirectory", staticmethod(lambda *a, **k: ""))

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)

        win._on_extract_zip_clicked()
        assert win._zip_extract_worker is None


class TestDossierDetailWindow:
    def test_dung_voi_dossier_thieu_du_lieu(self, qtbot, db):
        from app.ui.dossier_detail_window import DossierDetailWindow

        docs = DocumentRepository(db)
        dossiers = DossierRepository(db)
        dossier = Dossier(dossier_code="HS000001", reference="ABCD1234EF", meta_count=0, debit_count=0, vat_count=0)
        dossiers.save(dossier)

        win = DossierDetailWindow(dossier, docs, dossiers)
        qtbot.addWidget(win)
        assert win.windowTitle() == "HỒ SƠ HS000001"

    def test_luu_review_va_ghi_chu(self, qtbot, db):
        from app.ui.dossier_detail_window import DossierDetailWindow

        docs = DocumentRepository(db)
        dossiers = DossierRepository(db)
        dossier = Dossier(dossier_code="HS000001", reference="ABCD1234EF", meta_count=0, debit_count=0, vat_count=0)
        dossier_id = dossiers.save(dossier)

        win = DossierDetailWindow(dossier, docs, dossiers)
        qtbot.addWidget(win)
        win._notes_edit.setText("kiểm tra lại số tiền")
        win._reviewed_checkbox.setChecked(True)
        win._on_save()

        reloaded = dossiers.get(dossier_id)
        assert reloaded.notes == "kiểm tra lại số tiền"
        assert reloaded.reviewed is True
        assert reloaded.reviewed_at is not None

    def test_hien_thi_file_that_va_mo_pdf(self, qtbot, db):
        """Dùng PDF mẫu thật nếu có trong môi trường, bỏ qua nếu không."""
        sample = _first_sample_pdf()
        if sample is None:
            pytest.skip("Không có PDF mẫu trong môi trường này")

        from app.ui.dossier_detail_window import DossierDetailWindow

        docs_repo = DocumentRepository(db)
        dossiers_repo = DossierRepository(db)
        fields = FieldSet()
        fields.set(ExtractedField(field_name="invoice_number", value="1"))
        doc = Document(
            file_name=sample.name, file_path=sample, file_hash="h1",
            document_type=DocumentType.VPBANK_VAT_INVOICE, fields=fields,
        )
        doc_id = docs_repo.save(doc)

        dossier = Dossier(
            dossier_code="HS000001", reference="X", vat_document_id=doc_id,
            meta_count=0, debit_count=0, vat_count=1,
        )
        dossiers_repo.save(dossier)

        win = DossierDetailWindow(dossier, docs_repo, dossiers_repo)
        qtbot.addWidget(win)
        assert win.windowTitle() == "HỒ SƠ HS000001"


class TestStatusBadge:
    def test_doi_mau_theo_trang_thai(self, qtbot):
        from app.models.enums import DossierStatus
        from app.ui.view_models import StatusColor
        from app.ui.widgets.status_badge import StatusBadge

        badge = StatusBadge(DossierStatus.VALID)
        qtbot.addWidget(badge)
        assert StatusColor.GREEN in badge.styleSheet()
        badge.set_status(DossierStatus.MISSING_META)
        assert StatusColor.RED in badge.styleSheet()


class TestReviewWindow:
    def test_hien_thi_chung_tu_chua_ghep(self, qtbot, db):
        from app.ui.review_window import ReviewWindow

        docs_repo = DocumentRepository(db)
        fs = FieldSet()
        fs.set(ExtractedField(field_name="card_last4", value="1234"))
        doc = Document(
            file_name="a.pdf", file_path=Path("a.pdf"), file_hash="h1",
            document_type=DocumentType.META_INVOICE, fields=fs, raw_text="noi dung",
        )
        docs_repo.save(doc)

        win = ReviewWindow([doc], DossierRepository(db))
        qtbot.addWidget(win)
        assert win._unmatched_table.rowCount() == 1
        assert win._unmatched_table.item(0, 0).text() == "a.pdf"

    def test_ghep_tay_tao_dossier_needs_review(self, qtbot, db):
        from app.models.enums import DossierStatus, MatchSource
        from app.ui.review_window import ReviewWindow

        docs_repo = DocumentRepository(db)
        dossiers_repo = DossierRepository(db)

        fs1 = FieldSet()
        fs1.set(ExtractedField(field_name="card_last4", value="1234"))
        fs1.set(ExtractedField(field_name="document_date", value=date(2026, 8, 1)))
        fs2 = FieldSet()
        fs2.set(ExtractedField(field_name="card_last4", value="1234"))
        fs2.set(ExtractedField(field_name="transaction_date", value=date(2026, 8, 1)))

        m = Document(file_name="m.pdf", file_path=Path("m.pdf"), file_hash="h1",
                     document_type=DocumentType.META_INVOICE, fields=fs1)
        d = Document(file_name="d.pdf", file_path=Path("d.pdf"), file_hash="h2",
                     document_type=DocumentType.VPBANK_DEBIT_NOTE, fields=fs2)
        docs_repo.save(m)
        docs_repo.save(d)

        win = ReviewWindow([m, d], dossiers_repo, next_dossier_index=1)
        qtbot.addWidget(win)
        assert len(win._candidates) == 1

        win._create_manual_dossier(m, d)

        saved = [dossiers_repo.get(row["id"]) for row in db.connection.execute("SELECT id FROM dossiers")]
        assert len(saved) == 1
        assert saved[0].status is DossierStatus.NEEDS_REVIEW
        assert saved[0].match_source is MatchSource.MANUAL
        assert saved[0].reference is None  # không tự gán reference khi ghép tay


class TestSettingsWindow:
    def test_dung_va_doc_gia_tri_mac_dinh(self, qtbot, db, app_settings):
        from app.ui.settings_window import SettingsWindow

        win = SettingsWindow(SettingsRepository(db), app_settings)
        qtbot.addWidget(win)
        assert win._fields["processing.min_chars_per_page"].value() == app_settings.min_chars_per_page

    def test_luu_va_doc_lai(self, qtbot, db, app_settings):
        from app.ui.settings_window import SettingsWindow

        repo = SettingsRepository(db)
        win = SettingsWindow(repo, app_settings)
        qtbot.addWidget(win)
        win._fields["accounting.marketing_expense"].setText("6417")
        for key, widget in win._fields.items():
            repo.set(key, str(win._get_widget_value(widget)))
        assert repo.get("accounting.marketing_expense") == "6417"


class TestDocumentPreview:
    def test_render_pdf_that(self, qtbot):
        sample = _first_sample_pdf()
        if sample is None:
            pytest.skip("Không có PDF mẫu trong môi trường này")

        from app.ui.document_preview import DocumentPreviewDialog

        dlg = DocumentPreviewDialog(str(sample))
        qtbot.addWidget(dlg)
        assert dlg._page_count >= 1
        assert not dlg._image_label.pixmap().isNull()

    def test_file_khong_ton_tai_khong_crash(self, qtbot):
        from app.ui.document_preview import DocumentPreviewDialog

        dlg = DocumentPreviewDialog("/khong/ton/tai.pdf")
        qtbot.addWidget(dlg)
        assert "Không đọc được" in dlg._image_label.text()


class TestLuongLamViecDayDu:
    """Kiểm chứng toàn bộ GUI + worker thread + service trên PDF thật.

    Đây là bài test tích hợp quan trọng nhất của Phase 4: mô phỏng đúng thao
    tác người dùng (bấm SCAN, chờ worker thread chạy xong, bấm MATCH) chứ
    không gọi thẳng service — nếu dây nối signal/slot giữa MainWindow và
    ScanWorker/MatchWorker bị sai, bài test này sẽ treo hoặc fail.
    """

    def test_scan_va_match_tren_pdf_that(
        self, qtbot, db, classifier_config, extraction_config, app_settings, accounting_config, misa_config
    ):
        if not _SAMPLE_DIR or not Path(_SAMPLE_DIR).is_dir():
            pytest.skip("Không có thư mục PDF mẫu trong môi trường này (đặt ACCOUNTING_SAMPLE_DIR)")
        expected_count = sum(1 for _ in iter_pdf_files(Path(_SAMPLE_DIR)))

        from app.ui.main_window import MainWindow

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)
        win._input_folder_edit.setText(_SAMPLE_DIR)

        win._on_scan_clicked()
        assert win._scan_worker is not None
        with qtbot.waitSignal(win._scan_worker.finished_ok, timeout=10000):
            pass
        qtbot.waitUntil(lambda: not win._scan_worker.isRunning(), timeout=2000)

        assert f"Đã quét {expected_count} chứng từ" in win._summary_label.text()

        win._on_match_clicked()
        assert win._match_worker is not None
        with qtbot.waitSignal(win._match_worker.finished_ok, timeout=10000):
            pass
        qtbot.waitUntil(lambda: not win._match_worker.isRunning(), timeout=2000)

        assert win._dossier_table.rowCount() >= 1
        assert win._dossier_table.item(0, 0).text() == "HS000001"
        assert win._export_button.isEnabled() is True

    def test_organize_va_export_qua_nut_bam_that(
        self, qtbot, db, classifier_config, extraction_config, app_settings,
        accounting_config, misa_config, monkeypatch, tmp_path,
    ):
        """Bấm ③ ORGANIZE PDF và ④ EXPORT EXCEL thật — không gọi thẳng
        OrganizeService/ExportService — để bắt lỗi dây nối signal/slot hoặc
        threading giống cách test SCAN/MATCH đã từng bắt được bug sqlite3
        check_same_thread. Giả lập hộp thoại chọn file bằng monkeypatch vì
        môi trường offscreen không có ai bấm OK."""
        if not _SAMPLE_DIR or not Path(_SAMPLE_DIR).is_dir():
            pytest.skip("Không có thư mục PDF mẫu trong môi trường này (đặt ACCOUNTING_SAMPLE_DIR)")

        from PySide6.QtWidgets import QFileDialog, QInputDialog

        from app.ui.main_window import MainWindow

        win = MainWindow(db, classifier_config, extraction_config, app_settings, accounting_config, misa_config)
        qtbot.addWidget(win)
        win._input_folder_edit.setText(_SAMPLE_DIR)
        output_root = tmp_path / "OUTPUT"
        win._output_folder_edit.setText(str(output_root))

        win._on_scan_clicked()
        with qtbot.waitSignal(win._scan_worker.finished_ok, timeout=10000):
            pass
        qtbot.waitUntil(lambda: not win._scan_worker.isRunning(), timeout=2000)

        win._on_match_clicked()
        with qtbot.waitSignal(win._match_worker.finished_ok, timeout=10000):
            pass
        qtbot.waitUntil(lambda: not win._match_worker.isRunning(), timeout=2000)

        # --- ORGANIZE PDF (thư mục xuất đã điền sẵn -> không cần hộp thoại) ---
        win._on_organize_clicked()
        assert win._organize_worker is not None
        with qtbot.waitSignal(win._organize_worker.finished_ok, timeout=10000):
            pass
        qtbot.waitUntil(lambda: not win._organize_worker.isRunning(), timeout=2000)

        # Không giả định giá trị reference cụ thể của bộ PDF mẫu — thư mục
        # xuất ra tên gì cũng được, miễn đúng tiền tố mã hồ sơ và có đủ 3 file.
        candidates = list(output_root.glob("HS000001_*"))
        assert len(candidates) == 1, f"Không tìm thấy đúng 1 thư mục HS000001_*: {candidates}"
        dossier_folder = candidates[0]
        assert (dossier_folder / "01_META_INVOICE.pdf").is_file()
        assert (dossier_folder / "02_VPBANK_DEBIT_NOTE.pdf").is_file()
        assert (dossier_folder / "03_VPBANK_VAT_INVOICE.pdf").is_file()

        # --- EXPORT EXCEL (giả lập QInputDialog + QFileDialog.getSaveFileName) ---
        excel_path = tmp_path / "ACCOUNTING_RESULT.xlsx"
        monkeypatch.setattr(QInputDialog, "getInt", staticmethod(lambda *a, **k: (52601, True)))
        monkeypatch.setattr(
            QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(excel_path), ""))
        )

        win._on_export_clicked()
        assert win._export_worker is not None
        with qtbot.waitSignal(win._export_worker.finished_ok, timeout=10000):
            pass
        qtbot.waitUntil(lambda: not win._export_worker.isRunning(), timeout=2000)

        assert excel_path.is_file()
        import openpyxl

        wb = openpyxl.load_workbook(excel_path)
        misa_sheet = wb[misa_config.sheet_name]
        header = [c.value for c in misa_sheet[1]]
        so_ct_col = header.index("Số chứng từ (*)") + 1
        assert misa_sheet.cell(row=2, column=so_ct_col).value == "NVK052601"
