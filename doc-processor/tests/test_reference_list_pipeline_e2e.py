"""End-to-end: a reference-list Excel dropped into INPUT/ is picked up
automatically and cross-checked against the processed PDFs."""
from openpyxl import Workbook, load_workbook

from app.pipeline import run_pipeline
from tests.fixtures.generate_sample_pdfs import make_bank_debit_note, make_facebook_bill, make_vat_invoice


def _write_reference_list(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["Reference", "Amount"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_reference_list_flags_missing_reference_end_to_end(tmp_path):
    input_dir = tmp_path / "INPUT"
    fb_dir, vat_dir, bank_dir = input_dir / "facebook", input_dir / "vat", input_dir / "bank"
    for d in (fb_dir, vat_dir, bank_dir):
        d.mkdir(parents=True)

    make_facebook_bill(fb_dir / "fb.pdf", reference="LISTREF1", buyer_tax_id="0110382640")
    make_vat_invoice(
        vat_dir / "vat.pdf", vat_dir / "vat.xml", reference_code="LISTREF1",
        amount_before_vat="57.579", vat_amount="5.758", total_amount="63.337", buyer_tax_code="0110382640",
    )
    make_bank_debit_note(bank_dir / "bank.pdf", reference_code="LISTREF1")

    # reference list: one row matches the set above, one row has no PDFs at all
    _write_reference_list(
        input_dir / "expected_references.xlsx",
        [["LISTREF1", 63337], ["GHOSTREF9", 500000]],
    )

    output_dir = tmp_path / "OUTPUT"
    document_sets, summary = run_pipeline(input_dir, output_dir)

    by_ref = {ds.reference: ds for ds in document_sets}
    assert by_ref["LISTREF1"].reference_list_status == "FOUND"
    assert by_ref["LISTREF1"].category == "01_VALID"

    ghost = by_ref["GHOSTREF9"]
    assert ghost.reference_list_status == "NOT_FOUND"
    assert ghost.category == "09_MISSING_ALL_DOCUMENTS"

    ghost_folder = output_dir / "09_MISSING_ALL_DOCUMENTS" / "GHOSTREF9"
    assert (ghost_folder / "MISSING_01_Facebook.txt").exists()
    assert (ghost_folder / "MISSING_02_VAT_Invoice.txt").exists()
    assert (ghost_folder / "MISSING_03_Bank_Debit.txt").exists()

    wb = load_workbook(output_dir / "reconciliation_report.xlsx")
    assert "Reference List Check" in wb.sheetnames
    rows = {row[0].value: row[1].value for row in wb["Reference List Check"].iter_rows(min_row=2)}
    assert rows["LISTREF1"] == "FOUND"
    assert rows["GHOSTREF9"] == "NOT_FOUND"

    # the .xlsx reference list itself must never be treated as a PDF to process
    assert summary.total_files == 3


def test_no_reference_list_present_leaves_report_unchanged(tmp_path):
    input_dir = tmp_path / "INPUT" / "facebook"
    input_dir.mkdir(parents=True)
    make_facebook_bill(input_dir / "fb.pdf", reference="NOLIST1")

    output_dir = tmp_path / "OUTPUT"
    document_sets, _ = run_pipeline(input_dir.parent, output_dir)

    assert all(ds.reference_list_status is None for ds in document_sets)
    wb = load_workbook(output_dir / "reconciliation_report.xlsx")
    assert "Reference List Check" not in wb.sheetnames
