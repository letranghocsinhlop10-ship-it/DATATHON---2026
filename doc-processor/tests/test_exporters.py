from openpyxl import load_workbook

from app.export.misa_exporter import write_misa_excel
from app.export.reconciliation_report import write_reconciliation_report
from app.export.extracted_data_exporter import HEADERS, write_extracted_data_excel
from app.ingest import process_pdf_file
from app.matching.reference_matcher import match_documents
from app.organizing.categorizer import categorize
from app.organizing.file_organizer import organize_all
from app.validation.engine import validate_and_score
from tests.fixtures.generate_sample_pdfs import make_bank_debit_note, make_facebook_bill, make_vat_invoice


def _build_sets(tmp_path):
    # standard-format (non-dashed) tax code on both sides so this set comes
    # out fully VALID — the foreign-contractor-format nuance is already
    # covered by tests/test_validation.py.
    fb = process_pdf_file(
        make_facebook_bill(tmp_path / "fb.pdf", reference="REFAAA111", buyer_tax_id="0110382640")
    )
    vat_pdf, vat_xml = tmp_path / "vat.pdf", tmp_path / "vat.xml"
    make_vat_invoice(
        vat_pdf, vat_xml, reference_code="REFAAA111",
        amount_before_vat="57.579", vat_amount="5.758", total_amount="63.337",
    )
    vat = process_pdf_file(vat_pdf)
    bank = process_pdf_file(make_bank_debit_note(tmp_path / "bank.pdf", reference_code="REFAAA111"))

    fb2 = process_pdf_file(make_facebook_bill(tmp_path / "fb2.pdf", reference="REFBBB222"))

    sets = match_documents([fb, vat, bank, fb2])
    sets = [categorize(validate_and_score(ds)) for ds in sets]
    organize_all(sets, tmp_path / "OUTPUT")
    return sets


def test_misa_excel_has_configured_headers_and_rows(tmp_path):
    sets = _build_sets(tmp_path)
    out_path = tmp_path / "misa_import.xlsx"
    write_misa_excel(sets, out_path)

    wb = load_workbook(out_path)
    ws = wb["MISA Import"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == [
        "Reference", "Invoice Number", "Invoice Date", "Seller", "Tax Code", "Description",
        "Amount Before VAT", "VAT", "Total Amount", "Currency", "Payment Method",
        "Bank Transaction Date", "Document Status", "Check Result", "Issue", "Source Folder",
    ]
    references = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert "REFAAA111" in references
    assert "REFBBB222" in references

    complete_row = next(row for row in ws.iter_rows(min_row=2) if row[0].value == "REFAAA111")
    total_amount_col = headers.index("Total Amount")
    assert complete_row[total_amount_col].value == 63337.0


def test_reconciliation_report_has_expected_sheets_and_summary(tmp_path):
    sets = _build_sets(tmp_path)
    out_path = tmp_path / "reconciliation_report.xlsx"
    write_reconciliation_report(sets, out_path)

    wb = load_workbook(out_path)
    assert set(wb.sheetnames) == {"Summary", "Valid", "Errors", "Missing", "Duplicates", "All Sets", "MISA Import"}

    summary_rows = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary_rows["Tổng số bộ chứng từ"] == 2
    assert summary_rows["COMPLETE"] == 1
    assert summary_rows["INCOMPLETE"] == 1

    valid_refs = [row[0].value for row in wb["Valid"].iter_rows(min_row=2)]
    assert "REFAAA111" in valid_refs

    missing_refs = [row[0].value for row in wb["Missing"].iter_rows(min_row=2)]
    assert "REFBBB222" in missing_refs


def test_raw_extracted_excel_has_one_row_per_pdf(tmp_path):
    sets = _build_sets(tmp_path)
    out_path = tmp_path / "extracted_data.xlsx"
    write_extracted_data_excel(sets, out_path)

    wb = load_workbook(out_path)
    ws = wb["Extracted Data"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert headers == HEADERS
    assert ws.max_row == 5  # header + 3 PDFs in set A + 1 PDF in set B
