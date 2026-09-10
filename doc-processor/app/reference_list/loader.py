"""Loads a user-supplied Excel "reference list" — the set of reference
numbers (and optionally expected amount/invoice number/notes) an
accountant expects to see among the processed PDFs, e.g. exported from
a bank statement or an internal tracking sheet.

Column headers are matched flexibly (case/diacritic-insensitive against
a small alias list below) rather than requiring an exact header name,
since this file's format isn't standardized. Only a "reference" column
is required; the others are optional enrichments used for a light
cross-check against the extracted total_amount (see
app/reference_list/merge.py).
"""
from __future__ import annotations

import unicodedata
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from app.core.logging_config import get_logger
from app.extraction.base import normalize_amount
from app.models.schema import ReferenceListEntry

log = get_logger("reference_list.loader")

_HEADER_ALIASES: dict[str, list[str]] = {
    "reference": [
        "reference", "reference number", "ref", "so tham chieu", "ma tham chieu",
        "so gd", "ma giao dich", "transaction code",
    ],
    "expected_amount": [
        "amount", "expected amount", "so tien", "tong tien", "total amount", "total",
    ],
    "expected_invoice_number": [
        "invoice number", "invoice no", "so hoa don", "invoice", "hoa don",
    ],
    "notes": [
        "note", "notes", "ghi chu", "description", "dien giai", "noi dung",
    ],
}


def _normalize_header(s: str) -> str:
    """Lowercase, strip Vietnamese diacritics, collapse whitespace — so
    "Số tham chiếu", "SO THAM CHIEU" and "so_tham_chieu"-ish variants all
    match the same alias."""
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D")
    s = " ".join(s.lower().replace("_", " ").replace("-", " ").split())
    return s


def _match_columns(header_row: list[Optional[str]]) -> dict[str, int]:
    normalized = [_normalize_header(str(h)) if h else "" for h in header_row]
    columns: dict[str, int] = {}
    for field, aliases in _HEADER_ALIASES.items():
        for idx, header in enumerate(normalized):
            if header in aliases:
                columns[field] = idx
                break
    return columns


class ReferenceListFormatError(Exception):
    """Raised when the sheet has no recognizable 'reference' column."""


def load_reference_list(path: str | Path) -> list[ReferenceListEntry]:
    path = str(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]

    rows = ws.iter_rows(values_only=True)
    try:
        header_row = list(next(rows))
    except StopIteration:
        return []

    columns = _match_columns(header_row)
    if "reference" not in columns:
        raise ReferenceListFormatError(
            f"Không tìm thấy cột 'Reference/Số tham chiếu' trong {path}. "
            f"Các cột đọc được: {header_row}"
        )

    entries: list[ReferenceListEntry] = []
    for row_number, row in enumerate(rows, start=2):
        ref_val = row[columns["reference"]] if columns["reference"] < len(row) else None
        if ref_val is None or str(ref_val).strip() == "":
            continue

        def cell(field: str):
            idx = columns.get(field)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        amount_raw = cell("expected_amount")
        expected_amount = None
        if amount_raw is not None:
            expected_amount = normalize_amount(str(amount_raw))

        invoice_raw = cell("expected_invoice_number")
        notes_raw = cell("notes")

        entries.append(
            ReferenceListEntry(
                reference=str(ref_val).strip(),
                expected_amount=expected_amount,
                expected_invoice_number=str(invoice_raw).strip() if invoice_raw else None,
                notes=str(notes_raw).strip() if notes_raw else None,
                source_file=path,
                row_number=row_number,
            )
        )

    log.info("Đọc %d dòng reference từ %s", len(entries), path)
    return entries


def discover_reference_list_files(input_dir: str | Path) -> list[str]:
    """Any .xlsx/.xls found anywhere under input_dir is treated as a
    reference-list file — there's no other role for a spreadsheet in
    this pipeline's input."""
    input_dir = Path(input_dir)
    if not input_dir.is_dir():
        return []
    found = sorted(input_dir.rglob("*.xlsx")) + sorted(input_dir.rglob("*.xls"))
    # Skip Excel's own lock files (~$name.xlsx) left behind by an open file.
    return [str(p) for p in found if not p.name.startswith("~$")]


def load_all_reference_lists(paths: list[str]) -> list[ReferenceListEntry]:
    entries: list[ReferenceListEntry] = []
    for path in paths:
        try:
            entries.extend(load_reference_list(path))
        except ReferenceListFormatError as exc:
            log.warning("Bỏ qua file reference list không hợp lệ %s: %s", path, exc)
        except Exception:
            log.exception("Lỗi khi đọc reference list %s", path)
    return entries
